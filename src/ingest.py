"""
Ingestione input multimodali per analisi funzionale automatica.
Prende screenshot, note testuali, HTML, PDF, DOCX ed estrae un contesto strutturato.
Gli screenshot vengono analizzati con Vision API dell'LLM.

Usage:
    python ingest.py --input-dir ./inputs --output-file project_context.md
    
Environment Variables:
    LLM_API_KEY: La tua chiave API (OBBLIGATORIA per analisi screenshot)
    LLM_PROVIDER: Provider (openai, anthropic, google, dashscope)
    LLM_BASE_URL: URL base dell'API (opzionale, override)
    LLM_MODEL: Modello da usare (opzionale, override)
"""

import os
import sys
import argparse
import base64
import json
from pathlib import Path
from datetime import datetime

from bs4 import BeautifulSoup

import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from config import LLM_CONFIG, DEFAULT_PROVIDER

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

DEFAULT_INPUT_DIR = "./inputs"
DEFAULT_OUTPUT_FILE = "output/context/project_context.md"

# ---------------------------------------------------------------------------
# Helper Functions
# ---------------------------------------------------------------------------

def read_text_files(input_dir: str) -> list[dict]:
    """Read all .txt and .md files from input directory."""
    texts = []
    input_path = Path(input_dir)
    
    for ext in ["*.txt", "*.md"]:
        for file_path in input_path.glob(ext):
            try:
                content = file_path.read_text(encoding="utf-8")
                texts.append({
                    "type": "text",
                    "filename": file_path.name,
                    "content": content
                })
                print(f"  Loaded text: {file_path.name}")
            except Exception as e:
                print(f"  Warning: Could not read {file_path.name}: {e}")
    
    return texts


def extract_html_structure(html_content: str, filename: str) -> dict:
    """Extract semantic structure from HTML file."""
    soup = BeautifulSoup(html_content, "html.parser")
    
    # Remove script and style tags
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    
    extracted = {
        "filename": filename,
        "title": soup.title.string if soup.title else "Untitled",
        "forms": [],
        "buttons": [],
        "inputs": [],
        "links": [],
        "headings": [],
        "semantic_structure": []
    }
    
    # Extract forms
    for form in soup.find_all("form"):
        form_data = {
            "action": form.get("action", ""),
            "method": form.get("method", "GET"),
            "fields": []
        }
        for field in form.find_all(["input", "select", "textarea"]):
            field_info = {
                "type": field.get("type", "text"),
                "name": field.get("name", ""),
                "id": field.get("id", ""),
                "required": field.has_attr("required"),
                "pattern": field.get("pattern", ""),
                "placeholder": field.get("placeholder", ""),
            }
            form_data["fields"].append(field_info)
        extracted["forms"].append(form_data)
    
    # Extract buttons
    for btn in soup.find_all(["button", "input[type='button']", "input[type='submit']"]):
        extracted["buttons"].append({
            "text": btn.get_text(strip=True),
            "type": btn.get("type", "button"),
            "id": btn.get("id", ""),
            "onclick": btn.get("onclick", ""),
            "formaction": btn.get("formaction", ""),
        })
    
    # Extract inputs (outside forms too)
    for inp in soup.find_all("input"):
        extracted["inputs"].append({
            "type": inp.get("type", "text"),
            "name": inp.get("name", ""),
            "id": inp.get("id", ""),
            "required": inp.has_attr("required"),
            "pattern": inp.get("pattern", ""),
            "placeholder": inp.get("placeholder", ""),
            "value": inp.get("value", ""),
        })
    
    # Extract links
    for link in soup.find_all("a", href=True):
        extracted["links"].append({
            "text": link.get_text(strip=True),
            "href": link.get("href", ""),
            "target": link.get("target", ""),
        })
    
    # Extract headings
    for i in range(1, 7):
        for h in soup.find_all(f"h{i}"):
            extracted["headings"].append({
                "level": i,
                "text": h.get_text(strip=True),
                "id": h.get("id", ""),
            })
    
    # Build semantic structure (simplified DOM tree)
    def extract_semantic(tag):
        result = []
        for child in tag.children:
            if child.name:
                node = {
                    "tag": child.name,
                    "id": child.get("id", ""),
                    "class": child.get("class", []),
                    "text": child.get_text(strip=True)[:100] if child.get_text(strip=True) else "",
                }
                if node["text"] or node["id"] or node["class"]:
                    result.append(node)
                result.extend(extract_semantic(child))
        return result
    
    body = soup.body if soup.body else soup
    extracted["semantic_structure"] = extract_semantic(body)
    
    return extracted


def process_html_files(input_dir: str) -> list[dict]:
    """Process all HTML files and extract structure."""
    html_structures = []
    input_path = Path(input_dir)
    
    for ext in ["*.html", "*.htm"]:
        for file_path in input_path.glob(ext):
            try:
                content = file_path.read_text(encoding="utf-8")
                structure = extract_html_structure(content, file_path.name)
                html_structures.append({
                    "type": "html",
                    "filename": file_path.name,
                    "structure": structure
                })
                print(f"  Extracted HTML structure: {file_path.name}")
            except Exception as e:
                print(f"  Warning: Could not process {file_path.name}: {e}")
    
    return html_structures


def analyze_screenshot_with_vision(image_path: str, mime_type: str) -> str:
    """
    Analizza uno screenshot con Vision API dell'LLM.
    Estrae: CTA, campi input, stati di errore, layout, flussi visibili.
    
    Args:
        image_path: Path del file immagine
        mime_type: MIME type dell'immagine
        
    Returns:
        Analisi testuale dello screenshot
    """
    api_key = os.getenv("LLM_API_KEY", "")
    if not api_key:
        return "⚠️  LLM_API_KEY non settato - analisi screenshot non disponibile"
    
    provider = os.getenv("LLM_PROVIDER", DEFAULT_PROVIDER)
    
    if provider in LLM_CONFIG:
        base_url = os.getenv("LLM_BASE_URL", LLM_CONFIG[provider]["base_url"])
        model = os.getenv("LLM_MODEL", LLM_CONFIG[provider]["model"])
    else:
        base_url = os.getenv("LLM_BASE_URL")
        model = os.getenv("LLM_MODEL")
        if not base_url or not model:
            return f"⚠️  Provider '{provider}' non configurato per Vision"
    
    # Leggi e codifica immagine
    with open(image_path, "rb") as f:
        image_data = base64.b64encode(f.read()).decode("utf-8")
    
    prompt = """You are a UI/UX Analyst. Analyze this screenshot and extract ALL functional information.

Extract and list:
1. **All visible UI elements**: buttons, inputs, links, forms, cards, lists, modals
2. **All Call-to-Action (CTA)**: what actions can the user take?
3. **All input fields**: what data can the user enter? (with labels, placeholders, validation hints)
4. **Error states visible**: are there any error messages, validation errors, empty states?
5. **Navigation elements**: menus, breadcrumbs, tabs, back buttons
6. **Layout structure**: what's the visual hierarchy? What's the primary action?
7. **Data displayed**: what information is shown? (lists, tables, charts, stats)
8. **User flow implied**: what is the user trying to accomplish on this screen?

Be SPECIFIC - list every element you see. Do NOT generalize.
Respond in Italian.
"""
    
    try:
        from openai import OpenAI
        client = OpenAI(api_key=api_key, base_url=base_url)
        
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": "Sei un UI/UX Analyst esperto. Analizza gli screenshot e estrai informazioni funzionali dettagliate."},
                {"role": "user", "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {
                        "url": f"data:{mime_type};base64,{image_data}",
                        "detail": "high"
                    }}
                ]}
            ],
            max_tokens=2048,
            temperature=0.3,
        )
        
        return response.choices[0].message.content.strip()
        
    except Exception as e:
        return f"⚠️  Errore analisi Vision: {e}"


def read_pdf_files(input_dir: str) -> list[dict]:
    """Read all PDF files and extract text content."""
    texts = []
    input_path = Path(input_dir)
    
    try:
        import pypdf
    except ImportError:
        print("  ⚠️  pypdf non installato - PDF non supportati. Installa: pip install pypdf")
        return texts
    
    for file_path in input_path.glob("*.pdf"):
        try:
            reader = pypdf.PdfReader(str(file_path))
            content = ""
            for page in reader.pages:
                content += page.extract_text() + "\n\n"
            
            texts.append({
                "type": "pdf",
                "filename": file_path.name,
                "content": content.strip()
            })
            print(f"  Loaded PDF: {file_path.name} ({len(reader.pages)} pagine)")
        except Exception as e:
            print(f"  Warning: Could not read {file_path.name}: {e}")
    
    return texts


def read_docx_files(input_dir: str) -> list[dict]:
    """Read all DOCX files and extract text content."""
    texts = []
    input_path = Path(input_dir)
    
    try:
        import docx
    except ImportError:
        print("  ⚠️  python-docx non installato - DOCX non supportati. Installa: pip install python-docx")
        return texts
    
    for file_path in input_path.glob("*.docx"):
        try:
            doc = docx.Document(str(file_path))
            content = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            texts.append({
                "type": "docx",
                "filename": file_path.name,
                "content": content.strip()
            })
            print(f"  Loaded DOCX: {file_path.name}")
        except Exception as e:
            print(f"  Warning: Could not read {file_path.name}: {e}")
    
    return texts


def process_screenshots(input_dir: str, use_vision: bool = True) -> list[dict]:
    """
    Process screenshot files and analyze with Vision API.
    
    Args:
        input_dir: Directory containing screenshots
        use_vision: If True, analyze each screenshot with Vision API
        
    Returns:
        List of screenshot dicts with analysis results
    """
    screenshots = []
    input_path = Path(input_dir)
    
    for ext in ["*.png", "*.jpg", "*.jpeg", "*.webp"]:
        for file_path in input_path.glob(ext):
            try:
                # Determina mime type dal suffisso reale del file (non dal glob pattern)
                suffix = file_path.suffix.lstrip(".").lower()
                if suffix == "jpg" or suffix == "jpeg":
                    mime_type = "image/jpeg"
                elif suffix == "png":
                    mime_type = "image/png"
                elif suffix == "webp":
                    mime_type = "image/webp"
                else:
                    mime_type = f"image/{suffix}"
                
                screenshot = {
                    "type": "screenshot",
                    "filename": file_path.name,
                    "mime_type": mime_type,
                    "file_path": str(file_path),
                    "analysis": None,
                }
                
                # Analyze with Vision if enabled and API key is set
                if use_vision and os.getenv("LLM_API_KEY"):
                    print(f"  🔍 Analyzing screenshot with Vision: {file_path.name}...")
                    analysis = analyze_screenshot_with_vision(str(file_path), mime_type)
                    screenshot["analysis"] = analysis
                else:
                    screenshot["analysis"] = "⚠️  Vision non disponibile (LLM_API_KEY non settato)"
                
                screenshots.append(screenshot)
                print(f"  ✓ Processed: {file_path.name}")
                
            except Exception as e:
                print(f"  Warning: Could not process {file_path.name}: {e}")
    
    return screenshots


def generate_context_markdown(texts: list, html_structures: list, screenshots: list, 
                               pdf_texts: list = None, docx_texts: list = None) -> str:
    """Generate the unified project context Markdown file."""
    
    # Merge all text sources
    all_texts = list(texts)  # copy
    if pdf_texts:
        all_texts.extend(pdf_texts)
    if docx_texts:
        all_texts.extend(docx_texts)
    
    sections = []
    
    # Header
    sections.append(f"""# Project Context
Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

This file contains the extracted context from all input files (notes, PDF, DOCX, screenshots, HTML).
It serves as the "Bible" for the functional analysis agents.

---

""")
    
    # Section 1: Business Rules (from text files, PDF, DOCX)
    sections.append("""## 1. Regole di Business (Estratte dalle note)

""")
    
    if all_texts:
        for text in all_texts:
            type_label = text["type"].upper()
            sections.append(f"### [{type_label}] {text['filename']}\n")
            sections.append(f"```\n{text['content']}\n```\n\n")
    else:
        sections.append("*Nessun file di testo trovato.*\n\n")
    
    sections.append("---\n\n")
    
    # Section 2: UI Inventory (from HTML)
    sections.append("""## 2. Inventario UI (Estratto da HTML)

""")
    
    if html_structures:
        for html in html_structures:
            structure = html["structure"]
            sections.append(f"### File: {structure['filename']}\n")
            sections.append(f"**Titolo:** {structure['title']}\n\n")
            
            if structure["forms"]:
                sections.append("#### Forms\n")
                for form in structure["forms"]:
                    sections.append(f"- **Action:** `{form['action']}` ({form['method']})\n")
                    for field in form["fields"]:
                        required = " (required)" if field["required"] else ""
                        sections.append(f"  - `{field['name']}`: {field['type']}{required}\n")
                        if field["pattern"]:
                            sections.append(f"    - Pattern: `{field['pattern']}`\n")
                sections.append("\n")
            
            if structure["buttons"]:
                sections.append("#### Buttons / CTA\n")
                for btn in structure["buttons"]:
                    sections.append(f"- `{btn['text']}` (type: {btn['type']})\n")
                    if btn.get("formaction"):
                        sections.append(f"  - Form action: `{btn['formaction']}`\n")
                sections.append("\n")
            
            if structure["inputs"]:
                sections.append("#### Input Fields\n")
                for inp in structure["inputs"]:
                    required = " (required)" if inp["required"] else ""
                    sections.append(f"- `{inp['name']}`: {inp['type']}{required}\n")
                sections.append("\n")
            
            if structure["links"]:
                sections.append("#### Links\n")
                for link in structure["links"]:
                    sections.append(f"- `{link['text']}` → `{link['href']}`\n")
                sections.append("\n")
            
            if structure["headings"]:
                sections.append("#### Headings Structure\n")
                for h in structure["headings"]:
                    indent = "  " * (h["level"] - 1)
                    sections.append(f"{indent}- H{h['level']}: {h['text']}\n")
                sections.append("\n")
            
            sections.append("---\n\n")
    else:
        sections.append("*Nessun file HTML trovato.*\n\n")
        sections.append("---\n\n")
    
    # Section 3: Screenshots (with Vision analysis)
    sections.append("""## 3. Screenshots UI (Analizzati con Vision)

""")
    
    if screenshots:
        for ss in screenshots:
            sections.append(f"### Screenshot: {ss['filename']}\n\n")
            if ss.get("analysis"):
                sections.append(f"{ss['analysis']}\n\n")
            else:
                sections.append("*Analisi non disponibile*\n\n")
            sections.append("---\n\n")
    else:
        sections.append("*Nessuno screenshot trovato.*\n\n")
    
    sections.append("---\n\n")
    
    # Section 4: Data Model Inference
    sections.append("""## 4. Modello Dati (Inferito)

Basato sull'analisi dei form e degli input HTML:

""")
    
    inferred_fields = {}
    for html in html_structures:
        for form in html["structure"]["forms"]:
            for field in form["fields"]:
                key = field["name"] or field["id"] or "unnamed"
                if key not in inferred_fields:
                    inferred_fields[key] = {
                        "types": set(),
                        "required": False,
                        "patterns": set(),
                    }
                inferred_fields[key]["types"].add(field["type"])
                if field["required"]:
                    inferred_fields[key]["required"] = True
                if field["pattern"]:
                    inferred_fields[key]["patterns"].add(field["pattern"])
    
    if inferred_fields:
        sections.append("| Campo | Tipo | Required | Pattern |\n")
        sections.append("|-------|------|----------|--------|\n")
        for field_name, info in sorted(inferred_fields.items()):
            types = ", ".join(sorted(info["types"]))
            patterns = ", ".join(sorted(info["patterns"])) if info["patterns"] else "-"
            required = "✓" if info["required"] else "-"
            sections.append(f"| `{field_name}` | {types} | {required} | {patterns} |\n")
        sections.append("\n")
    else:
        sections.append("*Nessun campo dati inferito dai form HTML.*\n\n")
    
    sections.append("---\n\n")
    
    # Section 5: API Endpoints (inferred)
    sections.append("""## 5. Endpoint API (Inferiti)

Basato sull'analisi delle action dei form:

""")
    
    endpoints = set()
    for html in html_structures:
        for form in html["structure"]["forms"]:
            if form["action"]:
                endpoints.add((form["action"], form["method"]))
        for btn in html["structure"]["buttons"]:
            if btn.get("formaction"):
                endpoints.add((btn["formaction"], "POST"))
    
    if endpoints:
        sections.append("| Endpoint | Method |\n")
        sections.append("|----------|--------|\n")
        for endpoint, method in sorted(endpoints):
            sections.append(f"| `{endpoint}` | {method} |\n")
        sections.append("\n")
    else:
        sections.append("*Nessun endpoint API inferito.*\n\n")
    
    sections.append("---\n\n")
    
    # Section 6: Notes for Agent
    sections.append("""## 6. Note per l'Agente Analista

Questo contesto è stato generato automaticamente. Utilizzalo come base per:

1. **Generare i flussi utente** (User Journey)
2. **Definire gli stati dell'applicazione** (State Machine)
3. **Identificare gli edge case** (gestione errori, stati limite)
4. **Creare diagrammi Mermaid** (Flowchart, Sequence Diagram)
5. **Generare configurazione XState** eseguibile

### Checklist per l'analisi:
- [ ] Ogni form ha uno stato di caricamento definito?
- [ ] Ogni chiamata API ha gestione errori 4xx e 5xx?
- [ ] C'è un modo per annullare ogni operazione intermedia?
- [ ] Cosa succede se l'utente perde la connessione?
- [ ] Cosa succede se l'utente preme "indietro" nel browser?
- [ ] Gli stati di errore mostrano messaggi chiari all'utente?
- [ ] C'è un modo per recuperare da uno stato di errore?

""")
    
    return "".join(sections)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Ingest multimodal inputs for functional analysis")
    parser.add_argument("--input-dir", type=str, default=DEFAULT_INPUT_DIR,
                        help=f"Directory containing input files (default: {DEFAULT_INPUT_DIR})")
    parser.add_argument("--output-file", type=str, default=DEFAULT_OUTPUT_FILE,
                        help=f"Output context file (default: {DEFAULT_OUTPUT_FILE})")
    args = parser.parse_args()
    
    input_path = Path(args.input_dir)
    if not input_path.exists():
        print(f"Creating input directory: {args.input_dir}")
        input_path.mkdir(parents=True, exist_ok=True)
        print("Please add your input files (text, HTML, screenshots) to this directory and re-run.")
        return
    
    print(f"Processing inputs from: {args.input_dir}")
    print()
    
    # Process all inputs
    print("Step 1: Reading text files...")
    texts = read_text_files(args.input_dir)
    print(f"  Found {len(texts)} text file(s)")
    print()
    
    print("Step 2: Reading PDF files...")
    pdf_texts = read_pdf_files(args.input_dir)
    print(f"  Found {len(pdf_texts)} PDF file(s)")
    print()
    
    print("Step 3: Reading DOCX files...")
    docx_texts = read_docx_files(args.input_dir)
    print(f"  Found {len(docx_texts)} DOCX file(s)")
    print()
    
    print("Step 4: Processing HTML files...")
    html_structures = process_html_files(args.input_dir)
    print(f"  Found {len(html_structures)} HTML file(s)")
    print()
    
    print("Step 5: Processing screenshots (with Vision analysis)...")
    screenshots = process_screenshots(args.input_dir, use_vision=True)
    print(f"  Found {len(screenshots)} screenshot(s)")
    print()
    
    # Generate context
    print("Step 6: Generating unified context...")
    context_md = generate_context_markdown(texts, html_structures, screenshots, pdf_texts, docx_texts)
    
    # Write output
    with open(args.output_file, "w", encoding="utf-8") as f:
        f.write(context_md)
    
    print(f"  Wrote context to: {args.output_file}")
    print()
    print("Done! Ready to run spec.py")


if __name__ == "__main__":
    main()